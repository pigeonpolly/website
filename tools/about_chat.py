#!/usr/bin/env python3
"""Вопросы чата «Поболтать с Алиной и Полли» (/about-me/, docs/site.md).

  tools/about_chat.py table.xlsx      — взять таблицу Алины (A вопрос, B ответ Алины, C ответ Полли, строка 1 — заголовки):
                                        порядок и список вопросов — как в таблице, ответы RU — оттуда; изменившийся текст стирает
                                        старые EN/LV — их надо перевести заново
  tools/about_chat.py --todo          — что ещё не переведено (JSON: id → {q|a|p: текст RU})
  tools/about_chat.py --apply tr.json — вписать переводы: {id: {q|a|p: {en, lv}}}
  tools/about_chat.py --xlsx out.xlsx — выгрузить таблицу с текущими вопросами и ответами RU

Данные — content/about-chat/questions.json, на сайт их кладёт build.py (site/assets/about-chat.json).
"""
import json
import sys
from pathlib import Path

FILE = Path(__file__).resolve().parent.parent / "content" / "about-chat" / "questions.json"
FIELDS = ("q", "a", "p")


def load():
    return json.loads(FILE.read_text())


def save(d):
    FILE.write_text(json.dumps(d, ensure_ascii=False, indent=1) + "\n")


def cell(v):
    return str(v).strip() if v is not None else ""


def set_ru(obj, text):
    """Новый текст RU; если он поменялся — старые переводы уже не годятся."""
    if obj.get("ru", "") == text:
        return obj
    return {"ru": text} if text else {}


def import_xlsx(path):
    import openpyxl
    d = load()
    by_ru = {x["q"].get("ru"): x for x in d["questions"]}
    n = max([int(x["id"][1:]) for x in d["questions"]] + [0])
    out = []
    for row in list(openpyxl.load_workbook(path, data_only=True).active.iter_rows(values_only=True))[1:]:
        q, a, p = (cell(row[i]) if i < len(row) else "" for i in range(3))
        if not q:
            continue
        x = by_ru.pop(q, None)
        if x is None:
            n += 1
            x = {"id": "q%03d" % n, "q": {"ru": q}, "a": {}, "p": {}}
            print("новый вопрос", x["id"], q)
        x["a"], x["p"] = set_ru(x["a"], a), set_ru(x["p"], p)
        out.append(x)
    for x in by_ru.values():
        print("убран", x["id"], x["q"].get("ru"))
    d["questions"] = out
    save(d)
    todo = sum(1 for x in out for f in FIELDS if x[f].get("ru") and not (x[f].get("en") and x[f].get("lv")))
    print(f"вопросов {len(out)}, ответов Алины {sum(1 for x in out if x['a'].get('ru'))}, Полли {sum(1 for x in out if x['p'].get('ru'))}, без перевода {todo}")


def todo():
    d = load()
    t = {}
    for x in d["questions"]:
        for f in FIELDS:
            if x[f].get("ru") and not (x[f].get("en") and x[f].get("lv")):
                t.setdefault(x["id"], {})[f] = x[f]["ru"]
    print(json.dumps(t, ensure_ascii=False, indent=1))


def apply(path):
    d = load()
    tr = json.loads(Path(path).read_text())
    by_id = {x["id"]: x for x in d["questions"]}
    for i, fs in tr.items():
        for f, langs in fs.items():
            by_id[i][f].update({k: v for k, v in langs.items() if k in ("en", "lv") and v})
    save(d)
    print("переводов вписано:", sum(len(v) for v in tr.values()))


def export(path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Вопросы"
    ws.append(["Вопрос", "Ответ Алины", "Ответ Полли"])
    for x in load()["questions"]:
        ws.append([x["q"].get("ru", ""), x["a"].get("ru", ""), x["p"].get("ru", "")])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="F0A987")
    for col, w in zip("ABC", (52, 70, 70)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a:
        print(__doc__)
    elif a[0] == "--todo":
        todo()
    elif a[0] == "--apply":
        apply(a[1])
    elif a[0] == "--xlsx":
        export(a[1])
    else:
        import_xlsx(a[0])
